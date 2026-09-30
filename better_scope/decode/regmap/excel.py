"""Excel register-map importer (one workbook per device).

Template (see ``docs/REGISTER_MAPS.md`` and ``resources/register_map_template.xlsx``):

- Optional sheet ``Device``: key/value rows -- ``name``, ``bus``
  (smbus/pmbus/spi), ``address`` (7-bit), ``addr_width``, ``data_width``.
- Sheet ``Registers`` (or the first other sheet): one row per field.
  Register-level columns (``Register``, ``Address``, ``Width``, ``Access``,
  ``Reset``) only need filling on a register's first row; later rows with
  both ``Register`` and ``Address`` empty belong to the same register.

Headers are matched case-insensitively through :data:`COLUMN_ALIASES`; a
parenthesised suffix is ignored (``Width (bits)`` = ``Width``). Every problem
found is collected and raised together as one
:class:`~better_scope.decode.regmap.model.RegmapError`.
"""

__lazy_modules__ = ["openpyxl", "openpyxl.styles", "openpyxl.utils"]

import datetime
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from better_scope.decode.regmap.importers import RegmapImporter
from better_scope.decode.regmap.model import BUS_TYPES, Device, Field, RegmapError, RegmapIssue, Register

# Canonical column -> accepted headers (already normalized, see _norm).
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "register": ("register", "reg", "register name", "reg name", "name"),
    "address": ("address", "addr", "offset", "reg addr", "register address", "reg address",
                "command", "command code", "cmd", "code"),
    "width": ("width", "size", "reg width", "register width", "bit width"),
    "access": ("access", "access type", "rw", "r/w", "type"),
    "reset": ("reset", "reset value", "default", "default value", "por"),
    "field": ("field", "field name", "bitfield", "bit field"),
    "bits": ("bits", "bit", "bit range", "range", "position"),
    "msb": ("msb", "bit msb", "high bit"),
    "lsb": ("lsb", "bit lsb", "low bit"),
    "description": ("description", "desc", "comment", "comments", "notes"),
    "enum": ("enum", "enums", "values", "enumeration", "encoding", "value map"),
}

# Device-sheet key -> accepted keys (normalized).
DEVICE_KEY_ALIASES: dict[str, tuple[str, ...]] = {
    "name": ("name", "device", "device name", "part", "part number"),
    "bus": ("bus", "bus type", "interface", "protocol"),
    "address": ("address", "addr", "device address", "i2c address", "smbus address", "slave address"),
    "addr_width": ("addr width", "address width", "register address width", "reg addr width"),
    "data_width": ("data width", "register width", "reg width", "data bits"),
}

TEMPLATE_HEADERS: tuple[str, ...] = (
    "Register", "Address", "Width", "Access", "Reset", "Field", "Bits", "Description", "Enum",
)

# Rows searched for the header row (real workbooks often have a title first).
HEADER_SEARCH_ROWS = 20

_REGISTER_LEVEL = ("register", "address", "width", "access", "reset")


def _norm(text: Any) -> str:
    """Normalize a header or key: drop ``(...)``, lower-case, ``_``/``-`` -> space."""
    s = re.sub(r"\(.*?\)", "", str(text)).replace("_", " ").replace("-", " ")
    return " ".join(s.split()).casefold()


def _lookup(aliases: dict[str, tuple[str, ...]]) -> dict[str, str]:
    return {alias: canonical for canonical, names in aliases.items() for alias in names}


_COLUMN_LOOKUP = _lookup(COLUMN_ALIASES)
_DEVICE_LOOKUP = _lookup(DEVICE_KEY_ALIASES)


# -- value parsers -----------------------------------------------------------


def parse_int(value: Any) -> int:
    """Parse a non-negative integer cell: ``0x1A``, ``1Ah``, ``0b11010``, ``26``.

    Raises:
        ValueError: If the value is not a non-negative integer in one of
            those forms.
    """
    if isinstance(value, bool):
        raise ValueError(f"expected a number, got {value!r}")
    if isinstance(value, int):
        out = value
    elif isinstance(value, float):
        if not value.is_integer():
            raise ValueError(f"expected an integer, got {value!r}")
        out = int(value)
    else:
        s = str(value).strip().replace("_", "")
        lowered = s.casefold()
        try:
            if lowered.startswith("0x"):
                out = int(s[2:], 16)
            elif lowered.startswith("0b"):
                out = int(s[2:], 2)
            elif lowered.endswith("h"):
                out = int(s[:-1], 16)
            else:
                out = int(s, 10)
        except ValueError:
            raise ValueError(f"cannot parse {value!r} as a number (use 0x1A, 1Ah, 0b11010 or 26)") from None
    if out < 0:
        raise ValueError(f"expected a non-negative number, got {value!r}")
    return out


def parse_bits(value: Any) -> tuple[int, int]:
    """Parse a bit range: ``7:4``, ``[7:4]``, ``7..4``, ``3`` or ``[3]``.

    Excel turns a typed ``7:4`` into the time 07:04; that is accepted too.

    Returns:
        ``(msb, lsb)`` with ``msb >= lsb``.

    Raises:
        ValueError: If the value is not a bit range.
    """
    if isinstance(value, datetime.time):
        hi, lo = value.hour, value.minute
    elif isinstance(value, (int, float)) and not isinstance(value, bool):
        hi = lo = parse_int(value)
    else:
        s = str(value).strip()
        if s.startswith("[") and s.endswith("]"):
            s = s[1:-1]
        parts = re.split(r"\s*(?::|\.\.)\s*", s.strip())
        if len(parts) not in (1, 2) or not all(p.isdigit() for p in parts):
            raise ValueError(f"cannot parse {value!r} as bits (use 7:4, [7:4], 7..4 or 3)")
        hi, lo = int(parts[0]), int(parts[-1])
    return max(hi, lo), min(hi, lo)


def parse_enum(value: Any) -> dict[int, str]:
    """Parse ``0=Off; 1=On; 2=Auto`` (``;`` or newline separated; ``:`` also works).

    Commas separate items only when there is no ``;`` or newline.

    Raises:
        ValueError: If an item has no ``=``/``:``, a bad key, or a duplicate key.
    """
    s = str(value).strip()
    items = re.split(r"[;\n]", s) if re.search(r"[;\n]", s) else s.split(",")
    out: dict[int, str] = {}
    for item in items:
        item = item.strip()
        if not item:
            continue
        m = re.match(r"^([^=:]+?)\s*[=:]\s*(.+)$", item)
        if not m:
            raise ValueError(f"enum item {item!r} is not 'value=label'")
        key = parse_int(m.group(1))
        if key in out:
            raise ValueError(f"enum value {key} appears twice")
        out[key] = m.group(2).strip()
    return out


# -- importer ----------------------------------------------------------------


@dataclass
class _PendingRegister:
    """A register being assembled from its rows."""

    row: int
    name: str
    address: int
    width: int | None = None
    access: str = ""
    reset: int | None = None
    description: str = ""
    fields: list[tuple[Field, int]] = field(default_factory=list)


class _SheetReader:
    """Reads one Registers sheet, collecting issues instead of stopping."""

    def __init__(self, sheet_name: str, rows: list[tuple[Any, ...]], device: Device, issues: list[RegmapIssue]) -> None:
        self.sheet = sheet_name
        self.rows = rows
        self.device = device
        self.issues = issues
        self.columns: dict[str, tuple[int, str]] = {}  # canonical -> (0-based index, header)

    def issue(self, message: str, row: int | None = None, column: str | None = None) -> None:
        col = None
        if column is not None and column in self.columns:
            index, header = self.columns[column]
            col = f"{get_column_letter(index + 1)} ({header})"
        self.issues.append(RegmapIssue(message, self.sheet, row, col))

    def cell(self, values: tuple[Any, ...], column: str) -> Any:
        """Cell value for a canonical column; blank strings become ``None``."""
        if column not in self.columns:
            return None
        index = self.columns[column][0]
        value = values[index] if index < len(values) else None
        if isinstance(value, str):
            value = value.strip() or None
        return value

    def find_header(self) -> int | None:
        """0-based index of the header row, mapping ``self.columns``."""
        for r, values in enumerate(self.rows[:HEADER_SEARCH_ROWS]):
            columns: dict[str, tuple[int, str]] = {}
            for c, value in enumerate(values):
                if value is None:
                    continue
                canonical = _COLUMN_LOOKUP.get(_norm(value))
                if canonical and canonical not in columns:
                    columns[canonical] = (c, str(value).strip())
            if "register" in columns and "address" in columns:
                self.columns = columns
                return r
        return None

    def read(self) -> None:
        header = self.find_header()
        if header is None:
            self.issue(f"no header row with Register and Address columns in the first {HEADER_SEARCH_ROWS} rows")
            return
        if "field" in self.columns and "bits" not in self.columns and not {"msb", "lsb"} <= self.columns.keys():
            self.issue("a Field column needs a Bits column or MSB and LSB columns", header + 1)
            return

        current: _PendingRegister | None = None
        broken = False  # the current register's first row was invalid; skip its rows
        for r in range(header + 1, len(self.rows)):
            values = self.rows[r]
            row = r + 1
            if all(self.cell(values, c) is None for c in self.columns):
                continue
            name = self.cell(values, "register")
            addr_raw = self.cell(values, "address")
            address: int | None = None
            if addr_raw is not None:
                try:
                    address = parse_int(addr_raw)
                except ValueError as e:
                    self.issue(str(e), row, "address")
                    self._finish(current)
                    current, broken = None, True
                    continue
            name = str(name) if name is not None else None

            continuation = (name is None and address is None) or (
                current is not None
                and (name is None or name == current.name)
                and (address is None or address == current.address)
            )
            if continuation:
                if current is None:
                    if not broken:
                        self.issue("field row before any register", row, "register")
                    continue
            else:
                self._finish(current)
                current, broken = None, False
                if name is None or address is None:
                    self.issue("a new register needs both Register and Address",
                               row, "register" if name is None else "address")
                    broken = True
                    continue
                current = _PendingRegister(row, name, address)
            self._register_level(current, values, row)
            self._field(current, values, row)
        self._finish(current)

    def _register_level(self, reg: _PendingRegister, values: tuple[Any, ...], row: int) -> None:
        """Take register-level cells (first non-blank value wins: forward fill)."""
        width = self.cell(values, "width")
        if width is not None and reg.width is None:
            try:
                reg.width = parse_int(width)
                if reg.width == 0:
                    raise ValueError("width must be at least 1 bit")
            except ValueError as e:
                self.issue(str(e), row, "width")
        access = self.cell(values, "access")
        if access is not None and not reg.access:
            reg.access = str(access).upper()
        reset = self.cell(values, "reset")
        if reset is not None and reg.reset is None:
            try:
                reg.reset = parse_int(reset)
            except ValueError as e:
                self.issue(str(e), row, "reset")
        description = self.cell(values, "description")
        if description is not None and self.cell(values, "field") is None and not reg.description:
            reg.description = str(description)

    def _field(self, reg: _PendingRegister, values: tuple[Any, ...], row: int) -> None:
        fname = self.cell(values, "field")
        bits = self.cell(values, "bits")
        msb, lsb = self.cell(values, "msb"), self.cell(values, "lsb")
        if fname is None:
            if bits is not None or msb is not None or lsb is not None:
                self.issue("bits given without a Field name", row, "field")
            return
        try:
            if bits is not None:
                hi, lo = parse_bits(bits)
            elif msb is not None or lsb is not None:
                hi = parse_int(msb if msb is not None else lsb)
                lo = parse_int(lsb if lsb is not None else msb)
                hi, lo = max(hi, lo), min(hi, lo)
            else:
                self.issue(f"field {fname!s} has no bits", row, "bits" if "bits" in self.columns else "msb")
                return
        except ValueError as e:
            self.issue(str(e), row, "bits" if bits is not None else "msb")
            return
        enum: dict[int, str] = {}
        enum_raw = self.cell(values, "enum")
        if enum_raw is not None:
            try:
                enum = parse_enum(enum_raw)
            except ValueError as e:
                self.issue(str(e), row, "enum")
            too_big = [k for k in enum if k >> (hi - lo + 1)]
            if too_big:
                self.issue(f"enum values {too_big} do not fit in field {fname!s} ({hi - lo + 1} bits)", row, "enum")
        description = self.cell(values, "description")
        reg.fields.append((Field(str(fname), lo, hi - lo + 1, str(description or ""), enum), row))

    def _finish(self, reg: _PendingRegister | None) -> None:
        """Validate a complete register and add it to the device."""
        if reg is None:
            return
        width = reg.width or self.device.data_width
        ok = True
        if reg.address >> self.device.addr_width:
            self.issue(f"address 0x{reg.address:X} does not fit in {self.device.addr_width} address bits",
                       reg.row, "address")
            ok = False
        if reg.address in self.device.registers:
            other = self.device.registers[reg.address]
            self.issue(f"duplicate address 0x{reg.address:X} (also register {other.name})", reg.row, "address")
            ok = False
        if reg.reset is not None and reg.reset >> width:
            self.issue(f"reset 0x{reg.reset:X} does not fit in {width} bits", reg.row, "reset")
        placed: list[tuple[Field, int]] = []
        names: set[str] = set()
        for f, row in reg.fields:
            bits_col = "bits" if "bits" in self.columns else "msb"
            if f.msb >= width:
                self.issue(f"field {f.name} [{f.bits}] is outside the {width}-bit register {reg.name}", row, bits_col)
                ok = False
                continue
            clash = next((g for g, _ in placed if g.mask & f.mask), None)
            if clash is not None:
                self.issue(f"field {f.name} [{f.bits}] overlaps field {clash.name} [{clash.bits}]", row, bits_col)
                ok = False
                continue
            if f.name.casefold() in names:
                self.issue(f"field name {f.name} appears twice in register {reg.name}", row, "field")
                ok = False
            names.add(f.name.casefold())
            placed.append((f, row))
        if not ok:
            return
        fields = tuple(sorted((f for f, _ in placed), key=lambda f: f.msb, reverse=True))
        self.device.registers[reg.address] = Register(
            reg.name, reg.address, width, reg.access, reg.reset, reg.description, fields
        )


def _read_device_sheet(sheet_name: str, rows: list[tuple[Any, ...]], device: Device, issues: list[RegmapIssue]) -> None:
    """Apply the key/value rows of a ``Device`` sheet to ``device``."""
    for r, values in enumerate(rows):
        if len(values) < 2 or values[0] is None:
            continue
        key = _DEVICE_LOOKUP.get(_norm(values[0]))
        value = values[1]
        if key is None or value is None or (isinstance(value, str) and not value.strip()):
            continue
        row = r + 1
        try:
            if key == "name":
                device.name = str(value).strip()
            elif key == "bus":
                bus = str(value).strip().casefold()
                if bus not in BUS_TYPES:
                    raise ValueError(f"bus must be one of {list(BUS_TYPES)}, got {value!r}")
                device.bus = bus
            elif key == "address":
                address = parse_int(value)
                if address > 0x7F:
                    raise ValueError(f"address 0x{address:X} is not a 7-bit address "
                                     f"(8-bit form? use 0x{address >> 1:02X})")
                device.address = address
            else:
                width = parse_int(value)
                if width == 0:
                    raise ValueError(f"{key} must be at least 1")
                setattr(device, key, width)
        except ValueError as e:
            issues.append(RegmapIssue(str(e), sheet_name, row, "B"))


def read_workbook(path: Path) -> Device:
    """Load a register-map workbook.

    Args:
        path: ``.xlsx``/``.xlsm`` file.

    Returns:
        The device.

    Raises:
        RegmapError: With every problem found (sheet/row/column located).
    """
    path = Path(path)
    try:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as e:
        raise RegmapError(str(path), [RegmapIssue(f"cannot open workbook: {e}")]) from e
    try:
        sheets = {ws.title: [tuple(row) for row in ws.iter_rows(values_only=True)] for ws in wb.worksheets}
    finally:
        wb.close()

    issues: list[RegmapIssue] = []
    device = Device(name=path.stem, source=str(path))
    device_sheet = next((name for name in sheets if _norm(name) == "device"), None)
    if device_sheet is not None:
        _read_device_sheet(device_sheet, sheets[device_sheet], device, issues)
    reg_sheet = next((name for name in sheets if _norm(name) == "registers"), None)
    if reg_sheet is None:
        reg_sheet = next((name for name in sheets if name != device_sheet), None)
    if reg_sheet is None:
        issues.append(RegmapIssue("no Registers sheet"))
    elif not issues:
        # Register checks depend on the device widths; skip them if the Device sheet is bad.
        _SheetReader(reg_sheet, sheets[reg_sheet], device, issues).read()
    if issues:
        raise RegmapError(str(path), issues)
    return device


class ExcelImporter(RegmapImporter):
    """Excel workbook importer (default template, alias headers)."""

    id = "excel"
    name = "Excel workbook"
    extensions = (".xlsx", ".xlsm")

    def load(self, path: Path) -> Device:
        """Read ``path`` with :func:`read_workbook`."""
        return read_workbook(path)


# -- template ----------------------------------------------------------------

_TEMPLATE_DEVICE: tuple[tuple[str, Any], ...] = (
    ("name", "EXAMPLE_PMIC"),
    ("bus", "smbus"),
    ("address", "0x40"),
    ("addr_width", 8),
    ("data_width", 8),
)

_TEMPLATE_ROWS: tuple[tuple[Any, ...], ...] = (
    ("CTRL", "0x01", 8, "RW", "0x00", "EN", "7", "Converter enable", "0=Off; 1=On"),
    (None, None, None, None, None, "SS", "4:3", "Soft-start time", "0=1 ms; 1=2 ms; 2=4 ms; 3=8 ms"),
    ("STATUS", "0x10", 8, "RO", "0x00", "FAULT", "7", "Latched fault (write 1 to clear)", None),
    (None, None, None, None, None, "UVP", "3", "Under-voltage", None),
    (None, None, None, None, None, "OCP", "2", "Over-current", None),
    (None, None, None, None, None, "OTP", "1", "Over-temperature", None),
    (None, None, None, None, None, "PGOOD", "0", "Power good", None),
    ("VOUT_CFG", "0x1C", 8, "RW", "0x0C", "EN", "7", "Output enable", "0=Off; 1=On"),
    (None, None, None, None, None, "MODE", "6:5", "Switching mode", "0=Off; 1=PWM; 2=Auto"),
    (None, None, None, None, None, "VSEL", "4:0", "Output voltage select", None),
    ("VOUT", "0x21", 16, "RW", "0x0000", "RANGE", "15:14", "Output range", "0=0.5-1 V; 1=1-2 V; 2=2-5 V"),
    (None, None, None, None, None, "CODE", "11:0", "Output voltage code", None),
    ("DEVICE_ID", "0x30", 8, "RO", "0x5A", None, None, "Device identifier (no fields)", None),
)

_TEMPLATE_NOTES: tuple[str, ...] = (
    "Better_Scope register map template -- see docs/REGISTER_MAPS.md.",
    "",
    "Device sheet (optional): name, bus (smbus / pmbus / spi), address (7-bit), addr_width, data_width.",
    "Registers sheet: one row per field. Register, Address, Width, Access and Reset only need",
    "filling on a register's first row; following rows with Register and Address empty are its fields.",
    "Address / Reset: 0x1A, 1Ah, 0b11010 or 26.  Width: bits (defaults to data_width).",
    "Bits: 7:4, [7:4], 7..4 or 3 (or separate MSB and LSB columns).",
    "Enum: 0=Off; 1=On; 2=Auto",
    "Bits not covered by any field are shown as RESERVED[msb:lsb] when non-zero.",
)


def write_template(path: Path) -> Path:
    """Write the example register-map workbook.

    Args:
        path: Destination ``.xlsx`` (parent folders are created).

    Returns:
        The path written.
    """
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    bold = Font(bold=True)

    ws_dev = wb.active
    ws_dev.title = "Device"
    ws_dev.append(("Key", "Value"))
    for row in _TEMPLATE_DEVICE:
        ws_dev.append(row)
    for cell in ws_dev[1]:
        cell.font = bold
    ws_dev.column_dimensions["A"].width = 14
    ws_dev.column_dimensions["B"].width = 18

    ws_reg = wb.create_sheet("Registers")
    ws_reg.append(TEMPLATE_HEADERS)
    for row in _TEMPLATE_ROWS:
        ws_reg.append(row)
    for cell in ws_reg[1]:
        cell.font = bold
    widths = (14, 10, 8, 8, 10, 10, 8, 34, 34)
    for index, width in enumerate(widths, start=1):
        ws_reg.column_dimensions[get_column_letter(index)].width = width
    # Keep typed bit ranges as text so Excel does not turn 7:4 into a time.
    bits_index = TEMPLATE_HEADERS.index("Bits") + 1
    for (cell,) in ws_reg.iter_rows(min_col=bits_index, max_col=bits_index, max_row=200):
        cell.number_format = "@"
    ws_reg.freeze_panes = "A2"

    ws_notes = wb.create_sheet("Notes")
    for line in _TEMPLATE_NOTES:
        ws_notes.append((line,))
    ws_notes.column_dimensions["A"].width = 100

    wb.save(path)
    return path
